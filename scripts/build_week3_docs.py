"""Build the Week 3 submission documents with the bundled artifact runtime.

Run with ~/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3.12.
The authoring-operation marker is a session responsibility and is not repeated here.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime'
DEFAULT_SKILL = Path.home() / '.codex/plugins/cache/openai-primary-runtime/documents/26.909.12148/skills/documents'
WRITEUP = 'Engineering Issue Desk Week 3 Writeup.docx'
DEMO = 'Engineering Issue Desk Week 3 Demo Guide.docx'
BLACK = RGBColor(0, 0, 0)


def paragraph(doc, text='', *, bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level)


def bullet(doc, text):
    return paragraph(doc, text, style='List Bullet')


def new_page(doc, title):
    p = heading(doc, title)
    p.paragraph_format.page_break_before = True


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = label
    for values in rows:
        cells = t.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.text = str(value)
    for row_index, row in enumerate(t.rows):
        row_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        row_pr.append(no_split)
        if row_index == 0:
            repeated = OxmlElement('w:tblHeader')
            row_pr.append(repeated)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                item = OxmlElement('w:' + edge)
                item.set(qn('w:val'), 'single')
                item.set(qn('w:sz'), '4')
                item.set(qn('w:color'), 'D9D9D9')
                borders.append(item)
            tc_pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', '95'), ('bottom', '95'), ('left', '110'), ('right', '110')]:
                item = OxmlElement('w:' + edge)
                item.set(qn('w:w'), value)
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            tc_pr.append(margins)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF1')
                tc_pr.append(shade)
            for p in cell.paragraphs:
                if headers[col_index] == 'Records':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.color.rgb = BLACK
                    r.bold = row_index == 0
    paragraph(doc).paragraph_format.space_after = Pt(1)
    return t


def base_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.72), Inches(.72)
    section.left_margin, section.right_margin = Inches(.78), Inches(.78)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = BLACK
        style.font.underline = False
        style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
        style._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Arial')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [('Title', 24, 0, 12), ('Subtitle', 11, 0, 14),
                                      ('Heading 1', 17, 12, 9), ('Heading 2', 12, 10, 5),
                                      ('Heading 3', 11, 8, 4)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = name != 'Subtitle'
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    doc.styles['List Bullet'].paragraph_format.space_after = Pt(5)
    doc.core_properties.author = 'Naresh Ganji'
    doc.core_properties.title = title
    doc.core_properties.subject = 'The Gen Academy Week 3 agent project'
    doc.core_properties.comments = ''
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(subtitle, 'Subtitle')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('Page ')
    run.font.name, run.font.size = 'Arial', Pt(9)
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    return doc


def architecture(root, qa, runtime):
    source = root / 'docs/week3-architecture.svg'
    source.write_text('''<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="340" viewBox="0 0 1200 340">
<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#333"/></marker></defs>
<rect width="1200" height="340" fill="white"/>
<g fill="#f3f5f6" stroke="#999" stroke-width="1.5"><rect x="25" y="20" width="330" height="90" rx="8"/><rect x="435" y="20" width="330" height="90" rx="8"/><rect x="845" y="20" width="330" height="90" rx="8"/><rect x="435" y="220" width="330" height="90" rx="8"/><rect x="845" y="220" width="330" height="90" rx="8"/></g>
<g fill="black" font-family="Arial,sans-serif" font-size="26" text-anchor="middle"><text x="190" y="58">Investigate</text><text x="190" y="89" font-size="20">Plan tools and specialists</text><text x="600" y="58">Validate citations</text><text x="600" y="89" font-size="20">Resolve saved quote IDs</text><text x="1010" y="58">Review evidence</text><text x="1010" y="89" font-size="20">Separate model support audit</text><text x="600" y="258">Human review</text><text x="600" y="289" font-size="20">Approve reject or revise</text><text x="1010" y="258">Finish</text><text x="1010" y="289" font-size="20">Save local decision</text></g>
<g fill="none" stroke="#333" stroke-width="2" marker-end="url(#arrow)"><path d="M355 65 H430"/><path d="M765 65 H840"/><path d="M1010 110 V160 H600 V215"/><path d="M765 265 H840"/><path d="M435 265 H190 V115"/></g>
<g fill="black" font-family="Arial,sans-serif" font-size="20"><text x="210" y="220">Revision with feedback</text><text x="210" y="246">and the prior report</text></g>
</svg>''')
    png = qa / 'architecture.png'
    js = 'const sharp=require(process.argv[1]); sharp(process.argv[2],{density:180}).png().toFile(process.argv[3]);'
    subprocess.run([str(runtime / 'dependencies/node/bin/node'), '-e', js,
                    str(runtime / 'dependencies/node/node_modules/sharp'), str(source), str(png)], check=True)
    return png


def read_metrics(path):
    return json.loads(path.read_text()) if path and path.exists() else None


def counted(value, singular):
    return f'{value:,} {singular}' + ('' if value == 1 else 's')


def metric_rows(metrics):
    if not metrics:
        return [('Final Week 3 evaluation', 'Awaiting recorded verification')]
    rows = []
    cases = metrics.get('cases', metrics.get('results', []))
    if isinstance(cases, dict):
        cases = [{'id': key, **value} for key, value in cases.items() if isinstance(value, dict)]
    for case in cases:
        if not isinstance(case, dict):
            continue
        label = str(case.get('id') or case.get('case_id') or case.get('scenario') or 'Recorded case')
        sample = case.get('sample', {}).get('path')
        if sample:
            label += ' ' + Path(sample).name
        status = case.get('status') or case.get('execution_status') or 'Recorded'
        counts = case.get('accounting', case.get('metrics', case))
        parts = ['Awaiting human review' if status == 'awaiting_review' else str(status)]
        for key, suffix in [('model_calls', 'model call'), ('tool_calls', 'tool call'),
                            ('total_tokens', 'reported token')]:
            value = counts.get(key)
            if key == 'total_tokens':
                value = counts.get('reported_token_usage', {}).get('total_tokens', value)
            if isinstance(value, (int, float)):
                parts.append(counted(value, suffix))
        if isinstance(case.get('report'), dict):
            report = case['report']
            reviewed = report.get('quality_review', {}).get('status') in {'completed', 'no_content'}
            label_kind = 'retained finding' if reviewed else 'draft finding'
            parts.append(counted(len(report.get('findings', [])), label_kind))
        rows.append((label[:50], '; '.join(parts)))
    return rows or [('Recorded evaluation', 'See evaluation/week3-results/report.json')]


def metric_summary(metrics):
    if not metrics:
        return 'Final Week 3 case results are awaiting recorded verification.'
    summary = metrics.get('summary', {})
    usage_note = ('Reported usage was available for every recorded model call. '
                  if summary.get('usage_complete') is True else
                  'Unavailable usage remains unknown; totals include only reported tokens. ')
    return (
        f'The recorded run scheduled {summary.get("scheduled_cases", "unreported")} cases and reached '
        f'{summary.get("completed_handoffs", "unreported")} saved human review handoffs, with '
        f'{counted(summary.get("failed_cases", 0), "failed case")}. It recorded '
        f'{summary.get("model_calls", "unreported")} model calls, '
        f'{summary.get("tool_calls", "unreported")} tool calls, and '
        f'{summary.get("reported_tokens", 0):,} reported tokens. Recorded AI review removals were '
        f'{counted(summary.get("ai_rejected_findings", 0), "finding")} and '
        f'{counted(summary.get("ai_rejected_next_steps", 0), "proposed next step")}. ' + usage_note +
        'A saved handoff is pending human judgment, not an approved answer or measured accuracy.'
    )


def build_writeup(args, diagram, metrics):
    doc = base_doc('Engineering Issue Desk Week 3 Project Writeup',
                   'Naresh Ganji\n' + args.verification_date)
    paragraph(doc, 'Engineering Issue Desk now supports a saved engineering investigation from intake through evidence review and a human decision. Deep Agents chooses research steps and can delegate; Fireworks supplies tool calling, Keenable supplies optional current web reads, and LangGraph preserves the workflow. The final result is a cited report that an engineer can reopen, challenge, revise, approve, or reject.')
    heading(doc, 'Project one liner')
    paragraph(doc, 'Engineering Issue Desk helps a Streamlit engineer investigate behavior and review code in a web workspace, replacing manual documentation searches and public issue comparison while researching with five domain read tools and bounded optional context, preserving evidence, and handing off before a final decision, with a proposed target of a usable investigation within ten minutes in eight of ten representative cases.')
    paragraph(doc, 'The time and usefulness target is proposed. Human usefulness, accuracy, citation entailment, and the manual workflow cost have not been measured. Recorded handoffs establish saved workflow completion on the tested cases.')
    heading(doc, 'User and scope')
    paragraph(doc, 'The user is a developer or engineering lead deciding what a Streamlit report means, which claims have support, and what focused checks to perform next. This independent use case adapts the handout code review pattern to engineering issue investigation. The original retrieval application remains available alongside the new agent.')
    bullet(doc, 'The agent can inspect supplied code, a 65-source snapshot, approved public pages, and bounded repository context when explicitly enabled.')
    bullet(doc, 'Specialists share one evidence registry and budget; static checks and historical review guidance have separate provenance and limits.')
    bullet(doc, 'Human approval is a local decision. A separate preview and confirmation workflow can prepare an approved report for a GitHub summary review; no live publication was performed.')
    paragraph(doc, 'The project does not execute supplied programs or tests, prove a private root cause, or treat an attributed public report as a verified diagnosis. The third code sample did not establish a defect. The final evaluation and earlier failures are reported separately.')
    paragraph(doc, 'Assignment: Week 3 Project Handout October 2026. Public code: https://github.com/naresh25ganji/engineering-issue-desk. Builder eligibility ends October 25 2026; certification ends November 11 2026. The user confirmed that no video has been recorded yet.')

    new_page(doc, 'Agent framework')
    fields = [
        ('Agent goal', 'Produce prioritized, cited findings, missing information, and focused proposed checks for one Streamlit engineering question. Findings remain unapproved until an engineer decides.'),
        ('Where people use it', 'The Deep Investigation tab and command line client use the same saved service. The Project guide tab exposes architecture, design, operations, validation, and submission assets.'),
        ('Steps in order', 'Validate intake and options, claim the case, investigate using tools and specialists, and validate saved citation IDs. Run an independent support review, pause for a human, then save a decision or investigate again with notes and the previous report.'),
        ('Available actions', 'Read the corpus, approved web pages, public issues and commit-pinned patches; optionally read related application code or prior reviewed guidance. Load maintained playbooks and run fixed static checks on registered code; external publication is a separate confirmed human action.'),
        ('What it remembers', 'Local SQLite preserves cases, immutable evidence, events, publication plans, and graph checkpoints across restarts. Opt-in cross-case recall uses bounded literal search of approved or rejected summaries; it is not semantic memory or Mem0.'),
        ('Hard limits', 'No host shell, supplied-program execution, arbitrary file read, automatic approval, or model-controlled publication. Source content and prior decisions remain untrusted; a prior approval cannot establish a new factual claim.'),
        ('Human involvement', 'An engineer inspects claims, exact quotes, provenance, and proposed checks before approving or rejecting. Up to two revisions require notes; a GitHub summary needs an approved case, exact preview confirmation, and unchanged PR commits.'),
        ('When a step fails', 'Tool errors remain visible and cannot become successful evidence. Explicit retry resumes an unfinished checkpoint; invalid completed reports can use at most two fresh repair conversations while preserving case limits.'),
        ('Success measure', 'The proposed measure is a human-judged usable handoff within ten minutes in eight of ten representative cases. Operational completion, literal quote integrity, AI verdicts, and human quality are separate measurements.'),
    ]
    for label, text in fields:
        heading(doc, label, 2)
        paragraph(doc, text)

    new_page(doc, 'Architecture and tool authority')
    paragraph(doc, 'An adaptive Deep Agents loop runs inside an application-controlled LangGraph. The model selects research and delegation; Python validates arguments, accounts for calls, owns immutable evidence, and controls the review transitions. Citation validation runs inside investigate, followed by separate evidence review and human review nodes.')
    doc.add_paragraph().add_run().add_picture(str(diagram), width=Inches(6.86))
    table(doc, ['Capability', 'Scope', 'Authority'], [
        ('Five domain reads', 'Snapshot search, Keenable search and fetch, public issue and PR patch reads', 'Bounded source evidence; no GitHub writes'),
        ('Related code option', 'app.py and triage_desk Python files; list, literal search, excerpt read', 'Default off and immutable per case'),
        ('Playbooks', 'Three maintained version 1.0.0 guidance bundles loaded on demand', 'Guidance, not citable source evidence'),
        ('Review history option', 'At most three prior approved or rejected prose summaries', 'Default off; revalidate with new evidence'),
        ('Bandit checks', 'Eight fixed rules on registered Python code only', 'Attributed static flags, not verified vulnerabilities'),
        ('Human decision', 'Approve, reject, or request revision', 'Approve/reject without a model; revision researches again'),
        ('PR summary preview', 'Approved public-source report; streamlit/streamlit PR only', 'Separate COMMENT publication only after exact confirmation'),
    ], [1.4, 3.0, 2.46])
    paragraph(doc, 'Deep Agents scratch files use a virtual StateBackend. The isolated analyzer is trusted fixed code; it parses saved source without running the supplied program. The investigator has no publication tool. Detailed component, trust-boundary, and sequence diagrams are available in the Project guide and repository docs.')

    new_page(doc, 'Datasets and provenance')
    paragraph(doc, 'The corpus contains 65 English records selected for Streamlit engineering questions. Its snapshot completed October 17 2026 Pacific and its API reference version is 1.66.0. On the October 25 submission date this is historical source material, not a claim that every record is current. The interface exposes age and version scope; refresh remains a maintainer action.')
    table(doc, ['Source type', 'Records', 'Use'], [
        ('Official documentation and API records', '38', 'Documented behavior and API context'),
        ('Public GitHub issue reports', '24', 'Attributed observations and reproduction details'),
        ('Contribution and template records', '2', 'Reporting guidance'),
        ('Bounded release excerpt', '1', 'Version context'),
        ('Total', '65', 'Curated allowlist with dates, hashes, pins, and licenses'),
    ], [3.0, .8, 3.06])
    paragraph(doc, 'The reproducible ingestion script normalizes Markdown, MDX, generated API docstrings, and public issue text while retaining technical prose, code, headings, and source links. It removes display wrappers, boilerplate, and obvious contact or credential strings. Full commit pins, source hashes, allowlists, and licenses are retained in data/manifest.json and data/licenses.')
    heading(doc, 'Evidence and evaluation inputs', 2)
    paragraph(doc, 'Keenable snippets and fetched pages retain their distinct evidence types and retrieval times. GitHub PR patches are tied to observed base and head commits. Supplied code and opt-in repository reads preserve exact text, paths or labels, line ranges, hashes, and timestamps; a later working-tree change creates another source snapshot.')
    paragraph(doc, 'The executable dataset evaluation/week3-cases.json contains three real code samples from reviews.py, freshness.py, and hosted_embeddings.py. Each exported run preserves original file and excerpt hashes and line ranges. Historical Week 2 retrieval experiments remain separate from Week 3 agent results.')
    heading(doc, 'Guidance is separate from evidence', 2)
    paragraph(doc, 'The three versioned playbooks cover Python security, test planning, and evidence causality. They are maintained tool-delivered guidance rather than native SkillsMiddleware. Optional review memory searches at most 200 recent approved or rejected cases and returns at most three bounded prose records, excluding pending/current cases, raw code, source bodies, and support quotes. Neither guidance source can support a factual citation.')

    new_page(doc, 'Prompts used during development')
    paragraph(doc, 'Naresh used Codex to extend the existing repository. The summaries below paraphrase the development direction. Coding instructions and runtime agent prompts have different roles.')
    heading(doc, 'User direction', 2)
    bullet(doc, 'Review the Week 3 handout and supplied solution kits, then expand Engineering Issue Desk with Deep Agents and the tools provided by the team.')
    bullet(doc, 'Use the confirmed Keenable and Fireworks access and start building the new investigation workflow.')
    bullet(doc, 'Add an architecture diagram, design documentation, more features, and validation of the working system.')
    bullet(doc, 'Review final submission readiness from all angles, include useful stretch options, and review the remaining video and submission work.')
    heading(doc, 'Implementation checks used with Codex', 2)
    paragraph(doc, 'Codex was directed to exercise the actual Deep Agents factory with controlled tool-calling models, preserve human review across restart, verify shared limits across specialists, and make provider failures inspectable. Later checks covered immutable repository consent, real process crashes, quote identity, reviewer schema failures, local history isolation, static analyzer bounds, and publication replay protection.')
    heading(doc, 'Runtime prompts', 2)
    paragraph(doc, 'The investigator prompt treats sources and code as untrusted data, requires registered evidence and quote IDs, attributes public reports, and separates proposed checks from observed results. It asks the agent to inspect relevant called helpers before asserting that a safeguard is missing, and to use playbooks or static analysis only within their scope.')
    paragraph(doc, 'The separate evidence reviewer receives selected findings, proposed steps, and selected quotes. Its prompt rejects unsupported causal remedies and assumptions about omitted code. It returns indexed verdicts only; application code filters originals and rejects incomplete or malformed coverage. Source locations are triage_desk/deep_investigation.py and triage_desk/agent_review.py.')
    paragraph(doc, 'The solution kits informed patterns, but the project adapts them to a distinct local investigation use case. Instructions inside supplied documents do not authorize external actions or replace the user request.')

    new_page(doc, 'Iterations and lessons')
    heading(doc, 'Evidence identity and support', 2)
    paragraph(doc, 'A live model modified quotation text and correctly failed literal validation. The report schema now accepts source-assigned quote IDs, and Python resolves the original text. An uploader workaround exposed a different problem: exact citations did not establish causal applicability. A separate AI audit was added, but it still retained advice that did not address the checkbox-triggered rerun. Reviewer instructions were strengthened; human quality remains unmeasured.')
    heading(doc, 'Durable execution and limits', 2)
    paragraph(doc, 'Process locks and startup recovery distinguish saved state from active ownership. Subprocess crash tests verify that OS lock release permits recovery and a checkpoint can resume without repeating a completed read. Invalid completed reports use fresh repair threads; revisions receive the previous report and notes. Structured outputs are stored as JSON dictionaries.')
    paragraph(doc, 'Automatic summarization could dispatch model calls outside the shared accounting path, so it was replaced by a deterministic context guard for every role. Dispatch is recorded before requests. One exploratory run reported 78,956 tokens after an in-flight call crossed the 64,000 pre-dispatch threshold; the threshold is not hard cancellation. Safe completion metadata preserves reported usage after truncated or malformed output.')
    heading(doc, 'Context and failure diagnosis', 2)
    paragraph(doc, 'The initial three-case run retained an unsupported missing-key claim when the configuration helper was outside the excerpt. Opt-in repository tools now support bounded discovery and explicit immutable reads. The later three-case run avoided that old claim but did not read the helper; it cannot be credited with a complete helper analysis. Bounded search can still miss later code.')
    paragraph(doc, 'Strict reviewer output rejected a real response. Finite safe failure codes now distinguish provider, context, schema, coverage, and explanation problems. Paired inline backticks around unchanged identifiers can be normalized without accepting links or unsafe markup. A later fresh run exposed a plain comparison expression incorrectly caught as HTML; the guard was corrected and regression tested. Review-only recovery resumed the saved case without repeating research.')
    heading(doc, 'Stretch features preserve authority', 2)
    paragraph(doc, 'Maintained playbooks provide procedures on demand. Local review history is an opt-in hint, not factual ground truth or Mem0. Bandit returns attributed flags from fixed static rules rather than a verified exploit. Publication requires a distinct approved-case preview and explicit confirmation; the agent cannot trigger it. These additions have separate checks; their availability does not imply use in the three-case evaluation.')

    new_page(doc, 'Verification and recorded results')
    paragraph(doc, f'The final offline regression suite passed {args.test_count} tests. Controlled models exercise the real Deep Agents tool loop, specialist evidence sharing, checkpoints, review, repairs, revisions, failure paths, and shared limits. Feature-specific checks cover repository boundaries, local guidance, the real static analyzer, and publication state guards. Offline behavior is separate from live answer quality.')
    paragraph(doc, metric_summary(metrics))
    paragraph(doc, 'The observed model roles were nine investigator calls and three independent reviews. Tool dispatches were two saved-evidence reads, one repository inventory, and three repository reads, including one rejected line range followed by smaller reads. No repository search, helper read, Keenable, delegation, playbook, history, or Bandit call occurred in this three-case run.')
    table(doc, ['Latest recorded case', 'Observed result'], metric_rows(metrics), [2.0, 4.86])
    paragraph(doc, 'The verified latest run retained two freshness findings in one case, while the other two cases retained no factual findings. Exact quote integrity passed for the nonempty report; it is not applicable to empty reports. The reviewer can be overconservative, so an empty result is not a demonstrated accuracy success. The freshness report recognizes that a timestamp equal to the current time is accepted.')
    paragraph(doc, 'The latest report file is ' + str(args.metrics_json.relative_to(ROOT) if args.metrics_json.is_relative_to(ROOT) else args.metrics_json) + '. Human quality fields remain unscored, and no evaluation case was automatically approved. Actual tool use and usage come from that report; availability does not mean a feature was invoked.')
    new_page(doc, 'Preserved runs and validation limits')
    table(doc, ['Preserved observation', 'Result and scope'], [
        ('Initial three-code run', '2 of 3 handoffs; 9 model calls, 3 saved-evidence reads, 44,382 reported tokens. All three drafts had exact citations, including the failed reviewer draft.'),
        ('Repository-context run', '3 of 3 handoffs; 12 model calls, 6 tools, 70,046 reported tokens. It listed/read repository files, but used no search, helper read, Keenable, or delegation.'),
        ('Separate checkpoint retry', 'One reviewer call, zero tools, 2,716 reported tokens; pending handoff, no repeated investigator call or approval.'),
        ('Separate extension exercise', 'One freshness case: 6 model calls, 4 tools, 34,120 reported tokens. Playbook loading and real Bandit ran; one finding and two test proposals remained pending review.'),
        ('Separate connectivity and reads', 'Fireworks tool round trip and five Keenable search results; direct snapshot, page, issue, and commit-pinned patch reads validated separately.'),
    ], [1.8, 5.06])
    paragraph(doc, 'The extension exercise opted into memory, but did not call history lookup; positive recall is covered offline. The historical runs remain unchanged. Several implementation and prompt conditions changed between runs, so their handoff counts do not isolate a causal quality improvement. The October 23 freshness report described a strictly earlier timestamp although equality is accepted; that wording was not carried forward as a verified conclusion. The hosted embedding sample did not establish a defect. AI filtering and exact citations do not replace human inspection.')

    paragraph(doc, 'A fresh October 25 run reached only two of three handoffs after valid comparison text was rejected as HTML. That failure remains exported in evaluation/week3-final-results. The corrected guard was regression tested; a separate review-only retry used one reviewer call, zero tools, and 3,144 reported tokens before the new verified run was recorded.')
    table(doc, ['Check', 'What it establishes', 'What remains unverified'], [
        ('603 offline tests', 'Controlled workflow and boundary behavior', 'General live answer quality'),
        ('Three live cases', 'Saved handoff completion and observed calls', 'Human usefulness or accuracy'),
        ('Extension exercise', 'Actual playbook and Bandit tool use', 'Live positive history recall'),
        ('55 publisher tests', 'Controlled target, revision, body, replay and HTTP guards', 'Real GitHub POST and write-token operation'),
    ], [1.3, 2.8, 2.76])
    paragraph(doc, 'The frozen source environment check found 117 packages current. Supported execution is from the repository checkout with uv; default wheel build discovery is not configured. The source ZIP includes maintained playbooks and a per-file hash manifest, while credentials, private cases, checkpoints, caches, and environments are excluded. Qualitative inspection notes are in evaluation/week3-verified-results/engineering-review.md; they are Codex-assisted observations, not user approval. A full database restore rehearsal and production deployment were not performed.')

    new_page(doc, 'Stretch options and submission readiness')
    paragraph(doc, 'The handout contains no official extra-credit feature list or scoring schedule. These are project stretch options. Their implemented scope and verification must be distinguished from untested external operations.')
    table(doc, ['Stretch option', 'Implemented boundary', 'Readiness evidence'], [
        ('On-demand playbooks', 'Three versioned and hashed maintained bundles; not native SkillsMiddleware or evidence', 'Catalog/tool tests and one separate observed live load'),
        ('Reviewed-case memory', 'Default-off, immutable opt-in; bounded local literal prose lookup', 'Isolation/filter tests; not Mem0 or semantic retrieval'),
        ('Static analysis', 'Real Bandit, eight fixed rules, registered code, isolated parser worker', 'Real analyzer tests and separate live call; human context required'),
        ('GitHub summary review', 'Approved-case preview, exact plan ID confirmation, current head/base check, atomic replay protection', 'Mocked HTTP/state tests only; no live POST or write token used'),
    ], [1.4, 3.0, 2.46])
    heading(doc, 'Remaining limits', 2)
    paragraph(doc, 'Each case shares 18 model calls, 28 tool calls, two revisions, and two report repairs. The 64,000 reported-token threshold and 180-second active-invocation deadline are checked before dispatch; in-flight calls may finish beyond them. Missing usage remains unknown. A deterministic 80,000-character context guard replaces automatic summarization.')
    paragraph(doc, 'This is a personal local workspace without hosted identity, tenant isolation, distributed execution, or encrypted case storage. Credential detection is heuristic. Source snapshots can be old and code excerpts incomplete. Static checks do not execute tests, and proposed fixes are not verified outcomes. Cross-case guidance is local literal recall; semantic memory and Mem0 remain unimplemented. Optional LangSmith configuration does not establish a live tracing validation.')
    heading(doc, 'Required handoff', 2)
    paragraph(doc, 'Submit a native Google Doc, a video of five minutes or less, and the public GitHub code link through https://forms.gle/waBYFkkf5xxmcouV9. The baseline Week 3 commit 6e87e4b is public and its GitHub CI passed. This final refresh must be pushed and its CI checked before calling the new version delivered. Native import and reviewer access must be confirmed separately.')
    paragraph(doc, 'The user confirmed that no video has been recorded yet. The companion guide targets 4 minutes 45 seconds, including a genuine live interaction and a visibly identified saved-case fallback. Video recording, access checks, and course-form submission remain explicit final tasks.')
    return doc


def cue(doc, timing, screen, narration):
    heading(doc, timing, 2)
    paragraph(doc, 'Screen cue: ' + screen, bold_lead='Screen cue:')
    paragraph(doc, 'Narration: ' + narration, bold_lead='Narration:')


def build_demo(args, metrics):
    doc = base_doc('Engineering Issue Desk Week 3 Demo Guide',
                   'Naresh Ganji\n' + args.verification_date + '\nFour minutes and forty five seconds')
    paragraph(doc, 'Use this guide to record the Week 3 demo. No video has been recorded yet. The sequence targets 4 minutes 45 seconds, leaving 15 seconds below the handout limit; demonstrate a genuine final-app interaction and identify any previously saved case visibly.')
    heading(doc, 'Prepare the recording')
    bullet(doc, 'Run the local app, check the microphone, and use readable zoom. Keep .env, account panels, private history, and credentials outside the recording surface.')
    bullet(doc, 'Open the current evaluation JSON by its explicit path and the Project guide diagram. The October 17 corpus is a dated snapshot; acknowledge any age warning.')
    bullet(doc, 'Choose a small real-code case with its source visible. Have a genuine saved case ID ready for latency fallback. Recheck its current report and review state before recording; do not assume a retained claim is correct.')
    bullet(doc, 'Leave review history off unless deliberately showing its consent and safe prose scope. Do not publish to GitHub during this demonstration; a local preview is enough to explain that separate capability.')
    heading(doc, 'Suggested input', 2)
    paragraph(doc, 'Statically review this small Python sample. Identify at most two claims supported by the registered code, distinguish observations from proposed checks, inspect relevant called helpers if related code is enabled, and pause for human review. Do not execute the sample or claim a reproduced failure.')
    paragraph(doc, 'Use the actual 30-line freshness.py sample selected by evaluation/week3-cases.json. The verified saved fallback is case 8c857321-9d0f-414f-a811-d291947bc391; it has observed playbook and Bandit calls and remains pending human review. Enable related application code only when relevant and understood. Never claim a tool ran until the timeline shows it.')
    cue(doc, 'Opening through 25 seconds', 'Show the app and Deep Investigation tab.',
        'Engineering Issue Desk helps an engineer investigate a Streamlit issue or review a small code sample. The agent chooses research steps and can delegate, while the application preserves evidence and pauses for a human decision.')
    cue(doc, '25 seconds through 1 minute 5 seconds', 'Supply the sample, show default-off options and start a real investigation. Keep progress visible.',
        'Fireworks powers the investigator. Keenable and public GitHub readers are available when research needs them. Related repository code and prior reviewed guidance require explicit opt-in. This run records its actual tool calls, sources, and usage; it does not execute the supplied program.')

    new_page(doc, 'Show tool use and evidence')
    cue(doc, '1 minute 5 seconds through 2 minutes 5 seconds', 'Show observed plan or tool events, then open the Project guide architecture while the run continues.',
        'Deep Agents can load maintained playbooks, inspect bounded sources, and ask specialists for help. Fixed Bandit checks can parse registered Python code, but a static flag is not a verified vulnerability. The roles share one budget and source registry. LangGraph checkpoints preserve unfinished work and the human review interrupt.')
    cue(doc, '2 minutes 5 seconds through 3 minutes 15 seconds', 'Open one report finding and its saved source, quote ID, line range/hash or dated provenance.',
        'The model selects a saved source and quote ID, and Python supplies the exact original passage. A separate evidence reviewer can remove unsupported content before the handoff. Both checks have limits: an exact quotation can support the wrong causal conclusion, and an AI reviewer can miss that error. I check what the code actually shows and what it leaves unknown.')
    heading(doc, 'If the live run is slow or fails', 2)
    paragraph(doc, 'Say: This new investigation is still running, or failed at the recorded step. I will visibly open a previously saved real case to demonstrate the final report and review controls. Its case ID and timeline identify the earlier run. Do not present its result as the new run or conceal the failure.')
    paragraph(doc, 'Explain that explicit retry resumes pending work without refunding calls. Invalid completed reports can use a fresh repair thread with the same evidence and allowance. Show only sanitized diagnostics; do not narrate raw provider output.')
    heading(doc, 'Stretch features in one sentence', 2)
    paragraph(doc, 'Playbooks supply maintained procedures, optional review memory supplies noncitable local guidance, Bandit supplies attributed static flags, and an approved-case publication preview requires another explicit confirmation. These are stretch options, not an official extra-credit checklist. Show only features actually present in the final build.')

    new_page(doc, 'Review and report the measured result')
    cue(doc, '3 minutes 15 seconds through 4 minutes 10 seconds', 'Inspect proposed checks and missing details, enter a genuine note, then choose a decision appropriate to the evidence. Reopen the case.',
        'These proposed tests have not been executed by the agent. I can approve the local report, reject it, or request a revision with notes and the prior report. Approval and rejection make no investigator call. A GitHub summary would need a separate preview, exact plan confirmation, and unchanged commits. This demonstration makes no external post. Saved checkpoints and optional local literal history are implemented; Mem0 is not.')
    cue(doc, '4 minutes 10 seconds through 4 minutes 45 seconds', 'Select the exact latest recorded evaluation path, then show the repository and final writeup title.',
        'Codex helped build and test the tool loop, state, and error recovery. Observed failures led to quote IDs, a separate evidence audit, and tighter context and budget rules. The recorded handoffs show workflow completion on these samples. Human usefulness and accuracy remain unmeasured, and the third sample did not establish a defect.')
    heading(doc, 'Numbers to read accurately', 2)
    summary = metrics['summary']
    paragraph(doc, f'The final offline suite passed {args.test_count} tests. The latest run saved {summary["completed_handoffs"]} of {summary["scheduled_cases"]} handoffs with {summary["failed_cases"]} failures, {summary["model_calls"]} model calls, {summary["tool_calls"]} tools, and {summary["reported_tokens"]:,} reported tokens with complete usage. AI review removed {summary["ai_rejected_findings"]} findings and {summary["ai_rejected_next_steps"]} proposed steps. Exact-citation checks apply to retained factual findings; empty reports are not applicable, not accuracy successes. Human quality remains unscored.')
    table(doc, ['Recorded case', 'Observed state'], metric_rows(metrics), [2.0, 4.86])
    heading(doc, 'After recording', 2)
    paragraph(doc, 'Check that the finished video is five minutes or less, readable, audible, and free of private content. Confirm the final native Google Doc and public repository open for reviewers, check CI for the pushed commit, and submit those links plus the video through the course form. Builder eligibility ends October 25 2026; certification ends November 11 2026. This guide does not mark recording or submission complete.')
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--metrics-json', type=Path, required=True)
    parser.add_argument('--test-count', type=int, required=True)
    parser.add_argument('--verification-date', default='October 25 2026')
    parser.add_argument('--runtime-root', type=Path, default=DEFAULT_RUNTIME)
    parser.add_argument('--skill-root', type=Path, default=DEFAULT_SKILL)
    parser.add_argument('--qa-dir', type=Path, default=ROOT / '.cache/week3-docs-qa')
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    runtime_python = args.runtime_root / 'dependencies/python/bin/python3.12'
    if Path(sys.executable).resolve() != runtime_python.resolve():
        raise SystemExit('Run this builder using the selected bundled artifact Python runtime.')
    args.qa_dir.mkdir(parents=True, exist_ok=True)
    (ROOT / 'docs').mkdir(exist_ok=True)
    diagram = architecture(ROOT, args.qa_dir, args.runtime_root)
    metrics = read_metrics(args.metrics_json)
    if not metrics or metrics.get('report_kind') != 'deep_agent_workflow':
        raise SystemExit('Supply an existing recorded Deep Agents workflow report.')
    if args.test_count < 1:
        raise SystemExit('Supply the observed positive final test count.')
    documents = [(WRITEUP, build_writeup(args, diagram, metrics), 'writeup'),
                 (DEMO, build_demo(args, metrics), 'demo')]
    sanitizer = args.skill_root / 'scripts/google_docs_title_sanitize.py'
    for filename, doc, label in documents:
        raw = args.qa_dir / ('raw-' + filename)
        destination = ROOT / 'docs' / filename
        doc.save(raw)
        subprocess.run([str(runtime_python), str(sanitizer), str(raw), '--out', str(destination)], check=True)
        subprocess.run([str(runtime_python), str(sanitizer), str(destination), '--check'], check=True)
        if args.render:
            render_dir = args.qa_dir / label
            render_dir.mkdir(exist_ok=True)
            for old_page in render_dir.glob('page-*.png'):
                old_page.unlink()
            subprocess.run([str(runtime_python), str(args.skill_root / 'render_docx.py'), str(destination),
                            '--output_dir', str(render_dir), '--emit_pdf'], check=True)
        print(destination)


if __name__ == '__main__':
    main()
